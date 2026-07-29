# -*- coding: utf-8 -*-
import openpyxl, json, importlib.util, io, contextlib
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
spec = importlib.util.spec_from_file_location("g", "grille436.py")
m = importlib.util.module_from_spec(spec)
with contextlib.redirect_stdout(io.StringIO()): spec.loader.exec_module(m)
GRID, JOURS = m.GRID, m.JOURS
C = json.load(open('cible436.json')); R = json.load(open('grille436.json'))
ORDER = sorted(C['cible'], key=lambda x: -R['slots'][x]); PER = ["matin", "après-midi"]
piv = {s: {j: {p: [] for p in PER} for j in JOURS} for s in C['cible']}
for t, ro, nj, sem in GRID:
    for j, day in zip(JOURS, sem):
        if day is None: continue
        for k, (sec, sav, prod) in enumerate(day): piv[sec][j][PER[k]].append((t, sav, prod))
def ent(sec,j,p,typ): return [(t, sav if typ=='sav' else prod) for t,sav,prod in piv[sec][j][p] if (sav if typ=='sav' else prod)]
def tot(sec,typ): return sum(n for j in JOURS for p in PER for _,n in ent(sec,j,p,typ))

F="Arial"; BLK=Font(name=F,size=10); BOLD=Font(name=F,size=10,bold=True)
WHT=Font(name=F,size=10,bold=True,color="FFFFFF"); SUB=Font(name=F,size=9,italic=True,color="595959")
TIT=Font(name=F,size=14,bold=True,color="1F3864"); H1=Font(name=F,size=9,bold=True,color="FFFFFF")
SMS=Font(name=F,size=9,bold=True,color="1A5AA8"); SMP=Font(name=F,size=9,bold=True,color="A04A18")
DIM=Font(name=F,size=10,color="808080"); GRN=Font(name=F,size=10,bold=True,color="0A7D0A")
FH=PatternFill("solid",fgColor="1F3864"); FH2=PatternFill("solid",fgColor="4D5D6A")
FSAV=PatternFill("solid",fgColor="E4EEFA"); FPROD=PatternFill("solid",fgColor="FCEEE6")
FS=PatternFill("solid",fgColor="EEF1EE"); FG=PatternFill("solid",fgColor="E2EFDA")
th=Side(style="thin",color="BFBFBF"); BOX=Border(th,th,th,th)
CC=Alignment(horizontal="center",vertical="center",wrap_text=True); LL=Alignment(horizontal="left",vertical="center")

wb = openpyxl.load_workbook('Planification_41_SAV_PROD.xlsx')
for nm in ("Planning SAV secteur","Planning PROD secteur","Contrôle secteur"):
    if nm in wb.sheetnames: del wb[nm]

def planning_sheet(typ, title, subtitle):
    ws = wb.create_sheet(title)
    ws["A1"]=title.upper(); ws["A1"].font=TIT
    ws["A2"]=subtitle; ws["A2"].font=SUB
    ws.column_dimensions['A'].width=24
    for i in range(2,15): ws.column_dimensions[get_column_letter(i)].width=11
    # en-têtes : jours sur 2 colonnes
    r=4
    c=ws.cell(row=r,column=1,value="Secteur"); c.font=H1; c.fill=FH; c.alignment=CC; c.border=BOX
    ws.merge_cells(start_row=r,start_column=1,end_row=r+1,end_column=1)
    for k,j in enumerate(JOURS):
        col=2+k*2
        ws.merge_cells(start_row=r,start_column=col,end_row=r,end_column=col+1)
        cc=ws.cell(row=r,column=col,value=j); cc.font=H1; cc.fill=FH; cc.alignment=CC
        for d in (0,1): ws.cell(row=r,column=col+d).border=BOX
        for d,lab in enumerate(("matin","après-midi")):
            c2=ws.cell(row=r+1,column=col+d,value=lab); c2.font=Font(name=F,size=8,color="FFFFFF")
            c2.fill=FH2; c2.alignment=CC; c2.border=BOX
    for d,lab in enumerate(("Ouvert","Demandé","Contrôle")):
        col=14+d
        ws.merge_cells(start_row=r,start_column=col,end_row=r+1,end_column=col)
        cc=ws.cell(row=r,column=col,value=lab); cc.font=H1; cc.fill=FH; cc.alignment=CC
        for rr in (r,r+1): ws.cell(row=rr,column=col).border=BOX
        ws.column_dimensions[get_column_letter(col)].width=11
    r+=2; st=r
    for s in ORDER:
        cc=ws.cell(row=r,column=1,value=s); cc.font=BOLD; cc.fill=FS; cc.alignment=LL; cc.border=BOX
        for k,j in enumerate(JOURS):
            for d,p in enumerate(PER):
                e=ent(s,j,p,typ); col=2+k*2+d
                cell=ws.cell(row=r,column=col)
                if e:
                    cell.value=" ".join(f"{t}·{n}" for t,n in e)
                    cell.font=SMS if typ=='sav' else SMP
                    cell.fill=FSAV if typ=='sav' else FPROD
                else:
                    cell.value="—"; cell.font=DIM
                cell.alignment=CC; cell.border=BOX
        dem = C['sav_dem'][s] if typ=='sav' else round(C['prod_dem'][s],1)
        ws.cell(row=r,column=14,value=tot(s,typ)).font=BOLD
        ws.cell(row=r,column=15,value=dem).font=DIM
        ws.cell(row=r,column=16,value=f'=IF($N{r}>=$O{r},"OK","NON")').font=BOLD
        ws.cell(row=r,column=16).fill=FG
        for col in (14,15,16): ws.cell(row=r,column=col).alignment=CC; ws.cell(row=r,column=col).border=BOX
        r+=1
    ws.cell(row=r,column=1,value="TOTAL")
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=13)
    ws.cell(row=r,column=2,value="12 secteurs · les totaux se recoupent avec l'onglet « Grille hebdo »")
    ws.cell(row=r,column=14,value=f"=SUM(N{st}:N{r-1})")
    ws.cell(row=r,column=15,value=f"=SUM(O{st}:O{r-1})")
    ws.cell(row=r,column=16,value=f'=COUNTIF(P{st}:P{r-1},"OK")&" / 12"')
    for col in range(1,17):
        c=ws.cell(row=r,column=col); c.font=WHT; c.fill=FH; c.alignment=CC; c.border=BOX
    ws.cell(row=r,column=2).alignment=LL
    ws.freeze_panes="B6"; ws.sheet_view.showGridLines=False
    ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    return ws

planning_sheet('sav',"Planning SAV secteur",
  "Créneaux réservés au SAV, ouverts à la prise de rendez-vous J+1/J+2. Format « T2·4 » = technicien T2, 4 créneaux. Chaque secteur doit apparaître 3 fois par semaine.")
planning_sheet('prod',"Planning PROD secteur",
  "Créneaux ouverts à la production, posés depuis le carnet J+5/J+7. Ils occupent tout ce que le SAV ne réserve pas.")

ws = wb.create_sheet("Contrôle secteur")
ws["A1"]="CONTRÔLE PAR SECTEUR"; ws["A1"].font=TIT
ws["A2"]="Trois conditions : au moins 3 passages SAV sur un rythme valide, réservation SAV ≥ demande, capacité PROD ≥ demande."
ws["A2"].font=SUB
cols=["Secteur","Jours de passage SAV","Passages","Rythme","SAV ouvert","SAV demandé","Marge SAV",
      "PROD ouverte","PROD demandée","Marge PROD","Total créneaux","Cible","Contrôle"]
w=[24,22,10,10,11,12,10,12,13,11,13,9,11]
for i,(lab,wd) in enumerate(zip(cols,w),1):
    c=ws.cell(row=4,column=i,value=lab); c.font=H1; c.fill=FH; c.alignment=CC; c.border=BOX
    ws.column_dimensions[get_column_letter(i)].width=wd
ws.row_dimensions[4].height=30
r=5; st=r
for s in ORDER:
    jp=R['pass'][s]
    ry=("L-Me-V" if set(jp)=={"Lundi","Mercredi","Vendredi"} else
        "Ma-J-S" if set(jp)=={"Mardi","Jeudi","Samedi"} else "quotidien")
    sv,pr=tot(s,'sav'),tot(s,'prod'); sd,pd=C['sav_dem'][s],round(C['prod_dem'][s],1)
    vals=[s," · ".join(x[:3] for x in jp),len(jp),ry,sv,sd,None,pr,pd,None,None,round(C['cible'][s]),None]
    for i,v in enumerate(vals,1):
        c=ws.cell(row=r,column=i)
        if i==7: c.value=f"=$E{r}-$F{r}"; c.font=GRN
        elif i==10: c.value=f"=$H{r}-$I{r}"; c.font=GRN
        elif i==11: c.value=f"=$E{r}+$H{r}"; c.font=BOLD
        elif i==13: c.value=f'=IF(AND($C{r}>=3,$E{r}>=$F{r},$H{r}>=$I{r}),"conforme","À CORRIGER")'; c.font=BOLD; c.fill=FG
        else: c.value=v; c.font=BOLD if i in (1,5,8) else (DIM if i in (6,9,12) else BLK)
        c.alignment=LL if i in (1,2) else CC; c.border=BOX
        if i in (9,10): c.number_format="0.0"
    r+=1
ws.cell(row=r,column=1,value="TOTAL")
for col,f in ((5,f"=SUM(E{st}:E{r-1})"),(6,f"=SUM(F{st}:F{r-1})"),(7,f"=SUM(G{st}:G{r-1})"),
              (8,f"=SUM(H{st}:H{r-1})"),(9,f"=SUM(I{st}:I{r-1})"),(10,f"=SUM(J{st}:J{r-1})"),
              (11,f"=SUM(K{st}:K{r-1})"),(12,f"=SUM(L{st}:L{r-1})"),
              (13,f'=COUNTIF(M{st}:M{r-1},"conforme")&" / 12"')):
    ws.cell(row=r,column=col,value=f)
for col in range(1,14):
    c=ws.cell(row=r,column=col); c.font=WHT; c.fill=FH; c.alignment=CC; c.border=BOX
    if col in (9,10): c.number_format="0.0"
ws.freeze_panes="B5"; ws.sheet_view.showGridLines=False
ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1
ws.sheet_properties.pageSetUpPr.fitToPage=True

order=['Synthèse','Grille hebdo','Planning SAV secteur','Planning PROD secteur','Contrôle secteur',
       'Contrôle secteurs','Effectif','Demande 41','Moteur dispo','Distances']
wb._sheets=[wb[n] for n in order if n in wb.sheetnames]
wb.save('Planification_41_SAV_PROD.xlsx')
print("onglets :", wb.sheetnames)
